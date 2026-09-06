"""
Meeting Document Manager - Generates Word documents from meeting audit logs.
"""
import os
from datetime import datetime, timezone
from pathlib import Path

# Try to import python-docx, but handle gracefully if not installed
try:
    from docx import Document
    from docx.shared import Inches, Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# File path for the generated document
BASE_DIR = Path(__file__).resolve().parent.parent
DOCX_FILE_PATH = BASE_DIR / "meeting_discussion_details.docx"

def generate_meeting_word_doc(session_id: str = None) -> str:
    """
    Generates a Word document from the meeting audit log.
    Returns the path to the generated document.
    """
    from core.database import get_meeting_history
    
    if not DOCX_AVAILABLE:
        # Create a simple text file as fallback
        fallback_path = str(DOCX_FILE_PATH).replace('.docx', '.txt')
        history = get_meeting_history(session_id=session_id)
        with open(fallback_path, 'w', encoding='utf-8') as f:
            f.write(f"Meeting Transcript - Session: {session_id}\n")
            f.write(f"Generated: {datetime.now(timezone.utc).isoformat()}\n\n")
            for msg in history:
                f.write(f"[{msg.get('timestamp', '')}] {msg.get('speaker_name', '')} ({msg.get('speaker_role', '')}): {msg.get('response_text', '')}\n\n")
        return fallback_path
    
    # Get meeting history
    if session_id:
        history = get_meeting_history(session_id=session_id)
    else:
        history = get_meeting_history()
    
    if not history:
        return str(DOCX_FILE_PATH)
    
    # Create document
    doc = Document()
    doc.add_heading('Meeting Discussion Details', 0)
    
    # Session info
    p = doc.add_paragraph()
    p.add_run(f"Session: ").bold = True
    p.add_run(session_id or "All Sessions")
    
    p = doc.add_paragraph()
    p.add_run(f"Generated: ").bold = True
    p.add_run(datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"))
    
    p = doc.add_paragraph()
    p.add_run(f"Total Messages: ").bold = True
    p.add_run(str(len(history)))
    
    doc.add_heading('Transcript', level=1)
    
    for msg in history:
        # Speaker info
        p = doc.add_paragraph()
        speaker_info = f"{msg.get('speaker_name', 'Unknown')} ({msg.get('speaker_role', 'Unknown')})"
        if msg.get('avatar'):
            speaker_info = f"{msg.get('avatar', '')} {speaker_info}"
        run = p.add_run(speaker_info)
        run.bold = True
        run.font.size = Pt(11)
        
        # Timestamp
        timestamp = msg.get('timestamp', '')
        if timestamp:
            p.add_run(f"  [{timestamp}]").italic = True
        
        # Prompt
        p = doc.add_paragraph()
        p.add_run("Prompt: ").bold = True
        p.add_run(msg.get('user_prompt', ''))
        
        # Response
        p = doc.add_paragraph()
        p.add_run("Response: ").bold = True
        p.add_run(msg.get('response_text', ''))
        
        # Separator
        doc.add_paragraph("—" * 40)
    
    # Save
    doc.save(str(DOCX_FILE_PATH))
    return str(DOCX_FILE_PATH)


if __name__ == "__main__":
    # Test
    from core.database import init_db
    init_db()
    path = generate_meeting_word_doc()
    print(f"Document generated: {path}")